import os
import sqlite3
import time
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from typing import List, Optional

import bcrypt
import jwt
from fastapi import Depends, FastAPI, Header, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

DB_PATH = os.environ.get("DB_PATH", "/app/data/kochatzor.db")
JWT_SECRET = os.environ.get("JWT_SECRET")
if not JWT_SECRET:
    raise RuntimeError("JWT_SECRET majburiy (.env fayliga qo'shing)")

TOKEN_TTL_SECONDS = 180 * 24 * 3600  # 180 kun

app = FastAPI(title="Ko'chatzor API")


@contextmanager
def get_conn():
    Path(DB_PATH).parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def init_db():
    with get_conn() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL,
                region_id INTEGER, region_name TEXT,
                district_id INTEGER, district_name TEXT,
                created_at INTEGER
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS surveys (
                id TEXT PRIMARY KEY,
                regionId INTEGER, districtId INTEGER, mahallaId INTEGER,
                region TEXT, district TEXT, mahalla TEXT,
                fio TEXT, phone TEXT, area REAL, tree TEXT, variety TEXT,
                count INTEGER, planting TEXT, source TEXT, payvandtag TEXT,
                createdAt INTEGER, updatedAt INTEGER, isDeleted INTEGER,
                submittedBy TEXT, syncedAt INTEGER
            )
            """
        )
        existing_cols = {row["name"] for row in conn.execute("PRAGMA table_info(surveys)")}
        for col in ("latitude", "longitude"):
            if col not in existing_cols:
                conn.execute(f"ALTER TABLE surveys ADD COLUMN {col} REAL NOT NULL DEFAULT 0")


init_db()


# ---- auth helpers ----
def make_token(user: sqlite3.Row) -> str:
    payload = {
        "sub": user["username"],
        "role": user["role"],
        "regionId": user["region_id"],
        "districtId": user["district_id"],
        "region": user["region_name"],
        "district": user["district_name"],
        "exp": int(time.time()) + TOKEN_TTL_SECONDS,
    }
    return jwt.encode(payload, JWT_SECRET, algorithm="HS256")


def current_user(authorization: Optional[str] = Header(default=None)) -> dict:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Avtorizatsiya talab qilinadi")
    token = authorization.split(" ", 1)[1]
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Token yaroqsiz yoki muddati tugagan")
    return payload


def require_admin(user: dict = Depends(current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Faqat admin uchun")
    return user


class LoginRequest(BaseModel):
    username: str
    password: str


@app.post("/api/auth/login")
def login(payload: LoginRequest):
    with get_conn() as conn:
        row = conn.execute(
            "SELECT * FROM users WHERE username = ?", (payload.username.strip().lower(),)
        ).fetchone()
    if not row or not bcrypt.checkpw(payload.password.encode(), row["password_hash"].encode()):
        raise HTTPException(status_code=401, detail="Login yoki parol noto'g'ri")
    return {
        "token": make_token(row),
        "role": row["role"],
        "regionId": row["region_id"],
        "districtId": row["district_id"],
        "region": row["region_name"],
        "district": row["district_name"],
    }


@app.get("/api/admin/users")
def list_users(user: dict = Depends(require_admin)):
    with get_conn() as conn:
        rows = conn.execute(
            "SELECT username, role, region_name, district_name FROM users "
            "ORDER BY region_name, district_name, username"
        ).fetchall()
    return [
        {
            "username": r["username"],
            "role": r["role"],
            "region": r["region_name"],
            "district": r["district_name"],
        }
        for r in rows
    ]


class PasswordChange(BaseModel):
    password: str


@app.put("/api/admin/users/{username}/password")
def change_password(username: str, payload: PasswordChange, user: dict = Depends(require_admin)):
    if len(payload.password) < 6:
        raise HTTPException(status_code=400, detail="Parol kamida 6 belgidan iborat bo'lishi kerak")
    pw_hash = bcrypt.hashpw(payload.password.encode(), bcrypt.gensalt()).decode()
    with get_conn() as conn:
        cur = conn.execute(
            "UPDATE users SET password_hash = ? WHERE username = ?",
            (pw_hash, username.strip().lower()),
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="Foydalanuvchi topilmadi")
    return {"ok": True}


@app.get("/api/health")
def health():
    return {"status": "ok"}


# ---- survey sync ----
class SurveyIn(BaseModel):
    id: str
    regionId: int
    districtId: int
    mahallaId: int
    region: str
    district: str
    mahalla: str
    fio: str
    phone: str
    area: float
    tree: str
    variety: str
    count: int
    planting: str
    source: str
    payvandtag: str = ""
    createdAt: int
    updatedAt: int
    isDeleted: bool = False
    latitude: float = 0.0
    longitude: float = 0.0


class SyncRequest(BaseModel):
    records: List[SurveyIn]


@app.post("/api/sync")
def sync(payload: SyncRequest, user: dict = Depends(current_user)):
    now = int(time.time() * 1000)
    is_admin = user["role"] == "admin"
    results = []
    with get_conn() as conn:
        for r in payload.records:
            # Dala xodimi faqat o'ziga biriktirilgan tumanga tegishli yozuv yubora oladi.
            region_id = r.regionId
            district_id = r.districtId
            region_name = r.region
            district_name = r.district
            if not is_admin:
                region_id = user["regionId"]
                district_id = user["districtId"]
                region_name = user["region"]
                district_name = user["district"]

            conn.execute(
                """
                INSERT INTO surveys (id, regionId, districtId, mahallaId, region, district,
                    mahalla, fio, phone, area, tree, variety, count, planting, source,
                    payvandtag, createdAt, updatedAt, isDeleted, submittedBy, syncedAt,
                    latitude, longitude)
                VALUES (:id, :regionId, :districtId, :mahallaId, :region, :district,
                    :mahalla, :fio, :phone, :area, :tree, :variety, :count, :planting, :source,
                    :payvandtag, :createdAt, :updatedAt, :isDeleted, :submittedBy, :syncedAt,
                    :latitude, :longitude)
                ON CONFLICT(id) DO UPDATE SET
                    regionId=excluded.regionId, districtId=excluded.districtId,
                    mahallaId=excluded.mahallaId, region=excluded.region,
                    district=excluded.district, mahalla=excluded.mahalla, fio=excluded.fio,
                    phone=excluded.phone, area=excluded.area, tree=excluded.tree,
                    variety=excluded.variety, count=excluded.count, planting=excluded.planting,
                    source=excluded.source, payvandtag=excluded.payvandtag,
                    createdAt=excluded.createdAt, updatedAt=excluded.updatedAt,
                    isDeleted=excluded.isDeleted, submittedBy=excluded.submittedBy,
                    syncedAt=excluded.syncedAt, latitude=excluded.latitude,
                    longitude=excluded.longitude
                """,
                {
                    "id": r.id, "regionId": region_id, "districtId": district_id,
                    "mahallaId": r.mahallaId, "region": region_name, "district": district_name,
                    "mahalla": r.mahalla, "fio": r.fio, "phone": r.phone, "area": r.area,
                    "tree": r.tree, "variety": r.variety, "count": r.count,
                    "planting": r.planting, "source": r.source, "payvandtag": r.payvandtag,
                    "createdAt": r.createdAt, "updatedAt": r.updatedAt,
                    "isDeleted": int(r.isDeleted), "submittedBy": user["sub"], "syncedAt": now,
                    "latitude": r.latitude, "longitude": r.longitude,
                },
            )
            results.append({"id": r.id, "updatedAt": r.updatedAt})
    return {"results": results}


def _filters_where(
    user: dict,
    q: Optional[str] = None,
    tree: Optional[str] = None,
    region: Optional[str] = None,
    district: Optional[str] = None,
    mahalla: Optional[str] = None,
    planting: Optional[str] = None,
    extra: str = "",
):
    where = "WHERE isDeleted=0" + extra
    params: list = []
    if user["role"] != "admin":
        where += " AND districtId = ?"
        params.append(user["districtId"])
    else:
        # Viloyat/tuman bo'yicha filtrlash faqat adminga kerak — oddiy xodim
        # allaqachon o'z tumaniga qulflangan.
        if region:
            where += " AND region = ?"
            params.append(region)
        if district:
            where += " AND district = ?"
            params.append(district)
    if mahalla:
        where += " AND mahalla = ?"
        params.append(mahalla)
    if planting:
        where += " AND planting = ?"
        params.append(planting)
    if tree:
        where += " AND tree = ?"
        params.append(tree)
    if q:
        where += " AND (fio LIKE ? OR phone LIKE ? OR mahalla LIKE ?)"
        params += [f"%{q}%", f"%{q}%", f"%{q}%"]
    return where, params


@app.get("/api/facets")
def facets(
    user: dict = Depends(current_user),
    region: Optional[str] = None,
    district: Optional[str] = None,
    mahalla: Optional[str] = None,
):
    # Kaskadli variantlar: viloyat tanlansa tumanlar shu viloyatnikiga, tuman
    # tanlansa MFY/mevalar/yillar shu tumannikiga qisqaradi va h.k.
    base_where = "WHERE isDeleted=0 AND trim(tree) != ''"
    base_params: list = []
    if user["role"] != "admin":
        base_where += " AND districtId = ?"
        base_params.append(user["districtId"])

    def distinct(col: str, where: str, params: list) -> list:
        rows = conn.execute(
            f"SELECT DISTINCT {col} FROM surveys {where} AND {col} != '' ORDER BY {col}",
            params,
        ).fetchall()
        return [r[0] for r in rows]

    with get_conn() as conn:
        district_where, district_params = base_where, list(base_params)
        if region:
            district_where += " AND region = ?"
            district_params.append(region)

        mahalla_where, mahalla_params = district_where, list(district_params)
        if district:
            mahalla_where += " AND district = ?"
            mahalla_params.append(district)

        scoped_where, scoped_params = mahalla_where, list(mahalla_params)
        if mahalla:
            scoped_where += " AND mahalla = ?"
            scoped_params.append(mahalla)

        result = {
            "mahallas": distinct("mahalla", mahalla_where, mahalla_params),
            "trees": distinct("tree", scoped_where, scoped_params),
            "years": distinct("planting", scoped_where, scoped_params),
        }
        if user["role"] == "admin":
            result["regions"] = distinct("region", base_where, base_params)
            result["districts"] = distinct("district", district_where, district_params)
    return result


@app.get("/api/records")
def list_records(
    user: dict = Depends(current_user),
    q: Optional[str] = None,
    tree: Optional[str] = None,
    region: Optional[str] = None,
    district: Optional[str] = None,
    mahalla: Optional[str] = None,
    planting: Optional[str] = None,
    limit: int = Query(default=200, le=1000),
    offset: int = 0,
):
    where, params = _filters_where(user, q, tree, region, district, mahalla, planting)
    with get_conn() as conn:
        rows = conn.execute(
            f"SELECT * FROM surveys {where} ORDER BY updatedAt DESC LIMIT ? OFFSET ?",
            params + [limit, offset],
        ).fetchall()
        total = conn.execute(f"SELECT COUNT(*) c FROM surveys {where}", params).fetchone()["c"]
    return {"total": total, "rows": [dict(row) for row in rows]}


class SurveyUpdate(BaseModel):
    fio: Optional[str] = None
    phone: Optional[str] = None
    area: Optional[float] = None
    tree: Optional[str] = None
    variety: Optional[str] = None
    count: Optional[int] = None
    planting: Optional[str] = None
    payvandtag: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None


@app.put("/api/records/{record_id}")
def update_record(record_id: str, payload: SurveyUpdate, user: dict = Depends(require_admin)):
    fields = {k: v for k, v in payload.model_dump().items() if v is not None}
    if not fields:
        raise HTTPException(status_code=400, detail="Yangilanadigan maydon yo'q")
    fields["updatedAt"] = int(time.time() * 1000)
    set_clause = ", ".join(f"{k} = :{k}" for k in fields)
    with get_conn() as conn:
        cur = conn.execute(
            f"UPDATE surveys SET {set_clause} WHERE id = :id",
            {**fields, "id": record_id},
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="Yozuv topilmadi")
    return {"ok": True}


@app.delete("/api/records/{record_id}")
def delete_record(record_id: str, user: dict = Depends(require_admin)):
    with get_conn() as conn:
        cur = conn.execute(
            "UPDATE surveys SET isDeleted=1, updatedAt=? WHERE id=?",
            (int(time.time() * 1000), record_id),
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="Yozuv topilmadi")
    return {"ok": True}


def _qr_payload(fio, tree, variety, payvandtag, planting) -> str:
    # Bir xil skaner bilan o'qilishi uchun Android ilovadagi Qr.kt bilan aynan bir xil format
    # (uz.kochatzor.util.Qr.payload): "KEY=value" qatorlari, "\n" bilan ajratilgan.
    def clean(v) -> str:
        return str(v if v is not None else "").replace("\n", " ").replace("\r", " ")

    return "\n".join([
        f"FIO={clean(fio)}",
        f"TUR={clean(tree)}",
        f"NAV={clean(variety)}",
        f"PAYVANDTAG={clean(payvandtag)}",
        f"YIL={clean(planting)[:4]}",
    ])


def _qr_png_bytes(payload: str) -> bytes:
    import qrcode
    from qrcode.constants import ERROR_CORRECT_M

    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_M, box_size=6, border=4)
    qr.add_data(payload)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@app.get("/api/export.xlsx")
def export_xlsx(
    user: dict = Depends(current_user),
    q: Optional[str] = None,
    tree: Optional[str] = None,
    region: Optional[str] = None,
    district: Optional[str] = None,
    mahalla: Optional[str] = None,
    planting: Optional[str] = None,
):
    import openpyxl
    from datetime import datetime, timedelta, timezone
    from itertools import groupby
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from PIL import Image as PILImage

    TASHKENT = timezone(timedelta(hours=5))

    def to_local_dt(epoch_ms):
        try:
            return datetime.fromtimestamp(epoch_ms / 1000, tz=TASHKENT).replace(tzinfo=None)
        except (TypeError, ValueError, OSError):
            return None

    # Faqat ko'chat biriktirilgan yozuvlar eksport qilinadi — ko'chatsiz (faqat
    # xonadon) yozuvlar android ilovaning o'z eksportida ham chiqarilmaydi.
    where, params = _filters_where(
        user, q=q, tree=tree, region=region, district=district, mahalla=mahalla,
        planting=planting, extra=" AND trim(tree) != ''",
    )
    # mahallaId+fio+phone bo'yicha ketma-ket guruhlanishi uchun (xonadon
    # ustunlarini merge qilishda qatorlar bir joyda turishi shart).
    query = (
        f"SELECT * FROM surveys {where} "
        "ORDER BY region, district, mahalla, fio, phone, createdAt"
    )
    with get_conn() as conn:
        rows = conn.execute(query, params).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ko'chatzor"
    headers = [
        "Viloyat", "Tuman", "MFY", "FIO", "Telefon", "Maydon (ga)", "Ko'chat turi",
        "Nav", "Soni", "Ekilgan sana", "Manba", "Payvandtag", "Kiritgan",
        "Kiritilgan vaqti", "QR kod",
    ]
    widths = [20, 18, 22, 24, 15, 12, 16, 16, 9, 13, 20, 14, 14, 17, 14]
    HOUSEHOLD_COLS = range(1, 7)  # Viloyat..Maydon — bitta xonadon uchun merge qilinadi

    border = Border(*(Side(style="thin", color="D0D7D3") for _ in range(4)))
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="006C4C")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    body_align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    zebra_fill = PatternFill("solid", fgColor="F2F8F4")
    CENTER_COLS = {5, 6, 9, 10, 14}  # Telefon, Maydon, Soni, Ekilgan sana, Kiritilgan vaqti (1-based)

    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = w
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    qr_col_idx = len(headers)
    qr_col = get_column_letter(qr_col_idx)
    ws.auto_filter.ref = f"A1:{qr_col}{len(rows) + 1}"

    row_i = 2
    group_key = lambda r: (r["mahallaId"], r["fio"], r["phone"])
    for group_idx, (_, group_rows) in enumerate(groupby(rows, key=group_key)):
        group_rows = list(group_rows)
        fill = zebra_fill if group_idx % 2 == 0 else None
        start_row = row_i
        for r in group_rows:
            household_values = [r["region"], r["district"], r["mahalla"], r["fio"], r["phone"], r["area"]]
            tree_values = [
                r["tree"], r["variety"], r["count"], r["planting"], r["source"],
                r["payvandtag"] or "-", r["submittedBy"] or "", to_local_dt(r["createdAt"]),
            ]
            # Xonadon ustunlari faqat guruhning birinchi qatorida yoziladi —
            # qolganlari merge qilingandan keyin bo'sh qoladi.
            values = (household_values if row_i == start_row else [None] * 6) + tree_values
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_i, column=col_idx, value=value)
                cell.border = border
                cell.alignment = body_align_center if col_idx in CENTER_COLS else body_align
                if fill:
                    cell.fill = fill
                if col_idx == 14:
                    cell.number_format = "yyyy-mm-dd hh:mm"
            ws.cell(row=row_i, column=qr_col_idx).border = border
            if fill:
                ws.cell(row=row_i, column=qr_col_idx).fill = fill

            payload = _qr_payload(r["fio"], r["tree"], r["variety"], r["payvandtag"], r["planting"])
            png_bytes = _qr_png_bytes(payload)
            xl_img = XLImage(PILImage.open(BytesIO(png_bytes)))
            xl_img.width = 90
            xl_img.height = 90
            ws.add_image(xl_img, f"{qr_col}{row_i}")
            ws.row_dimensions[row_i].height = 70
            row_i += 1

        end_row = row_i - 1
        if end_row > start_row:
            for col_idx in HOUSEHOLD_COLS:
                ws.merge_cells(start_row=start_row, start_column=col_idx, end_row=end_row, end_column=col_idx)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=kochatzor.xlsx"},
    )
